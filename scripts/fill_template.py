import json
from docx import Document

# Load generated resume
with open("./output/resume.json", "r", encoding="utf-8") as f:
    resume_data = json.load(f)

# Open template
doc = Document("./templates/TrentRice_Resume_Template.docx")


def replace_text(placeholder, value):
    for paragraph in doc.paragraphs:
        if placeholder in paragraph.text:
            paragraph.text = paragraph.text.replace(placeholder, value)


def replace_with_bullets(placeholder, text):
    for i, paragraph in enumerate(doc.paragraphs):
        if placeholder in paragraph.text:
            paragraph.text = ""

            bullets = text.split("\n")

            for bullet in bullets:
                bullet = bullet.replace("•", "").strip()

                if bullet:
                    new_para = paragraph.insert_paragraph_before(bullet)

                    new_para.style = "List Bullet"

            break


# Summary
replace_text("{{SUMMARY}}", resume_data.get("summary", ""))

# Experience Sections
replace_with_bullets("{{JDI_EXPERIENCE}}", resume_data.get("jdi_experience", ""))

replace_with_bullets("{{IMP_EXPERIENCE}}", resume_data.get("imp_experience", ""))

replace_with_bullets("{{TISSUE_EXPERIENCE}}", resume_data.get("tissue_experience", ""))

replace_with_bullets(
    "{{NSHEALTH_EXPERIENCE}}", resume_data.get("nshealth_experience", "")
)

replace_with_bullets(
    "{{SHIPBUILDING_EXPERIENCE}}", resume_data.get("shipbuilding_experience", "")
)


# Skills
replace_with_bullets(
    "{{PROGRAMMING_SKILLS}}", resume_data.get("programming_skills", "")
)
replace_with_bullets(
    "{{DATA_SCIENCE_SKILLS}}", resume_data.get("data_science_skills", "")
)
replace_with_bullets(
    "{{COMMUNICATION_SKILLS}}", resume_data.get("communication_skills", "")
)

# Save
doc.save("./output/tailored_resume.docx")

print("Resume created successfully.")
